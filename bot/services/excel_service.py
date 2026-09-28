from io import BytesIO
from datetime import datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from bot.database.models import FuelLog, GeneratorState, MaintenanceLog, RunLog, AuditResetLog
from bot.services.generator_service import format_dt, format_duration, utc_to_local


# Palette definition
NAVY_HEADER = "0F2942"
WHITE = "FFFFFF"
ZEBRA_EVEN = "F8FAFC"
ZEBRA_ODD = "FFFFFF"
BORDER_COLOR = "E2E8F0"
BORDER_DARK = "94A3B8"

# Accents
GREEN_BG = "DCFCE7"
GREEN_TXT = "166534"
RED_BG = "FEE2E2"
RED_TXT = "991B1B"
AMBER_BG = "FEF3C7"
AMBER_TXT = "92400E"
BLUE_BG = "E0F2FE"
BLUE_TXT = "0369A1"
PURPLE_BG = "EDE9FE"
PURPLE_TXT = "4338CA"

UKR_WEEKDAYS = ["Понеділок", "Вівторок", "Середа", "Четвер", "П'ятниця", "Субота", "Неділя"]


def apply_header_style(cell, text=None):
    """Applies modern deep navy styling to table header cells."""
    if text is not None:
        cell.value = text
    cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
    cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(border_style="thin", color=NAVY_HEADER)
    cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)


def apply_data_cell(cell, value, align="center", is_even=False, bold=False, text_color=None, bg_color=None, wrap_text=False):
    """Applies clean data formatting with optional zebra striping, alignment and text wrapping."""
    cell.value = value
    cell.font = Font(
        name="Calibri",
        size=10,
        bold=bold,
        color=text_color if text_color else "1E293B"
    )

    fill_hex = bg_color if bg_color else (ZEBRA_EVEN if is_even else ZEBRA_ODD)
    cell.fill = PatternFill(start_color=fill_hex, end_color=fill_hex, fill_type="solid")
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap_text)

    border_side = Side(border_style="thin", color=BORDER_COLOR)
    cell.border = Border(top=border_side, left=border_side, right=border_side, bottom=border_side)


def apply_total_row(ws, row_idx, num_cols, label="Всього:", sums=None):
    """Applies a professional double-underlined accounting total row at the bottom."""
    ws.row_dimensions[row_idx].height = 24
    for c in range(1, num_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        val = ""
        align = "center"
        if c == 1:
            val = label
            align = "left"
        elif sums and c in sums:
            val = sums[c]
            align = "right" if isinstance(val, (int, float)) or "₴" in str(val) else "center"

        cell.value = val
        cell.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        cell.alignment = Alignment(horizontal=align, vertical="center")

        top_border = Side(border_style="thin", color="64748B")
        bottom_border = Side(border_style="double", color=NAVY_HEADER)
        side_border = Side(border_style="thin", color=BORDER_COLOR)
        cell.border = Border(top=top_border, bottom=bottom_border, left=side_border, right=side_border)


def auto_fit_columns(ws, min_width=12, max_width=50):
    """Calculates optimal column widths based on contents and headers."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            for line in val_str.split("\n"):
                if len(line) > max_len:
                    max_len = len(line)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, max_width), min_width)


class ExcelService:

    @staticmethod
    async def generate_full_report(session: AsyncSession) -> BytesIO:
        wb = Workbook()

        # -----------------------------------------------------------------
        # 1. SUMMARY SHEET (📊 Зведення)
        # -----------------------------------------------------------------
        ws_summary = wb.active
        ws_summary.title = "📊 Зведення"
        ws_summary.sheet_properties.tabColor = "0284C7"
        ws_summary.views.sheetView[0].showGridLines = True
        ws_summary.freeze_panes = "A4"

        gen_res = await session.execute(select(GeneratorState).where(GeneratorState.id == 1))
        gen = gen_res.scalar_one_or_none()

        runs_res = await session.execute(select(RunLog).order_by(RunLog.id.desc()))
        runs = runs_res.scalars().all()

        fuel_res = await session.execute(select(FuelLog).order_by(FuelLog.id.desc()))
        refuels = fuel_res.scalars().all()

        maint_res = await session.execute(select(MaintenanceLog).order_by(MaintenanceLog.id.desc()))
        maints = maint_res.scalars().all()

        audit_res = await session.execute(select(AuditResetLog).order_by(AuditResetLog.id.desc()))
        audits = audit_res.scalars().all()

        total_hours_ran = sum(r.duration_hours for r in runs)
        total_fuel_burned = sum(r.fuel_consumed for r in runs)
        total_fuel_added = sum(f.amount_liters for f in refuels)
        total_fuel_cost = sum((f.cost or 0.0) for f in refuels)
        total_maint_cost = sum((m.cost or 0.0) for m in maints)

        # Main Title Banner
        ws_summary.merge_cells("A1:D1")
        title_cell = ws_summary["A1"]
        title_cell.value = f"⚡ ЗВІТ ДИСПЕТЧЕРА ГЕНЕРАТОРА: {gen.name if gen else 'Основний ДГУ'}"
        title_cell.font = Font(name="Calibri", size=15, bold=True, color=WHITE)
        title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_summary.row_dimensions[1].height = 36

        # Subtitle Row
        ws_summary.merge_cells("A2:D2")
        sub_cell = ws_summary["A2"]
        sub_cell.value = f"Сформовано: {format_dt(datetime.now(timezone.utc))} (Київський час) | MTservice Smart Generator Management"
        sub_cell.font = Font(name="Calibri", size=9, italic=True, color="64748B")
        sub_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        sub_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws_summary.row_dimensions[2].height = 20

        sections = [
            ("📌 ОСНОВНІ ПАРАМЕТРИ ОБ'ЄКТА", BLUE_BG, BLUE_TXT, [
                ("Поточний робочий стан", "🟢 В РОБОТІ" if (gen and gen.is_running) else "🔴 ЗУПИНЕНО"),
                ("Паспортна норма витрати пального", f"{gen.fuel_rate:.2f} л/год" if gen else "4.50 л/год"),
                ("Загальна ємність паливного бака", f"{gen.tank_capacity:.0f} л" if gen else "150 л"),
                ("Поточний залишок у баку", f"{gen.current_fuel:.1f} л ({(gen.current_fuel / gen.tank_capacity * 100):.0f}%)" if (gen and gen.tank_capacity) else "—"),
                ("Прогноз автономної роботи на залишку", f"~{(gen.current_fuel / gen.fuel_rate):.1f} год" if (gen and gen.fuel_rate) else "—"),
            ]),
            ("⏱ МОТОГОДИНИ ТА ТЕХНІЧНЕ ОБСЛУГОВУВАННЯ", PURPLE_BG, PURPLE_TXT, [
                ("Загальне напрацювання генератора", f"{gen.total_hours:.2f} мч" if gen else "0.00 мч"),
                ("Регламентний інтервал між ТО", f"{gen.maintenance_interval_hours:.0f} мч" if gen else "250 мч"),
                ("Позначка попереднього ТО (олива)", f"{gen.last_maintenance_hours:.1f} мч" if gen else "0.0 мч"),
                ("Залишок до наступного ТО", f"{((gen.last_maintenance_hours + gen.maintenance_interval_hours) - gen.total_hours):.2f} мч" if gen else "—"),
                ("Напрацювання свічок запалювання", f"{(gen.total_hours - (gen.last_spark_plugs_hours or 0.0)):.1f} мч тому" if gen else "—"),
                ("Напрацювання повітряного фільтру", f"{(gen.total_hours - (gen.last_air_filter_hours or 0.0)):.1f} мч тому" if gen else "—"),
                ("Напрацювання паливного фільтру", f"{(gen.total_hours - (gen.last_fuel_filter_hours or 0.0)):.1f} мч тому" if gen else "—"),
            ]),
            ("⛽ ПАЛИВНИЙ БАЛАНС ТА СТАТИСТИКА РОБОТИ", AMBER_BG, AMBER_TXT, [
                ("Всього зафіксовано запусків", f"{len(runs)} сесій"),
                ("Сумарний час роботи за весь період", f"{total_hours_ran:.2f} год ({format_duration(total_hours_ran)})"),
                ("Сумарна витрата пального за період", f"{total_fuel_burned:.2f} л"),
                ("Сумарно заправлено пального за період", f"{total_fuel_added:.2f} л"),
                ("Кількість проведених заправок", f"{len(refuels)} заправок"),
                ("Кількість виконаних ТО", f"{len(maints)} записів"),
            ]),
            ("💰 ФІНАНСОВІ ВИТРАТИ", GREEN_BG, GREEN_TXT, [
                ("Сумарні витрати на заправку пального", f"{total_fuel_cost:,.2f} ₴".replace(",", " ")),
                ("Сумарні витрати на сервіс та запчастини", f"{total_maint_cost:,.2f} ₴".replace(",", " ")),
                ("ЗАГАЛЬНИЙ БЮДЖЕТ ЕКСПЛУАТАЦІЇ", f"{(total_fuel_cost + total_maint_cost):,.2f} ₴".replace(",", " ")),
            ]),
        ]

        curr_row = 4
        for sec_title, sec_bg, sec_txt, items in sections:
            ws_summary.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=4)
            sec_cell = ws_summary.cell(row=curr_row, column=1)
            sec_cell.value = sec_title
            sec_cell.font = Font(name="Calibri", size=11, bold=True, color=sec_txt)
            sec_cell.fill = PatternFill(start_color=sec_bg, end_color=sec_bg, fill_type="solid")
            sec_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws_summary.row_dimensions[curr_row].height = 24
            curr_row += 1

            for idx, (param, val) in enumerate(items):
                ws_summary.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=2)
                p_cell = ws_summary.cell(row=curr_row, column=1, value=param)
                p_cell.font = Font(name="Calibri", size=10, bold=False, color="334155")
                p_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

                ws_summary.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=4)
                v_cell = ws_summary.cell(row=curr_row, column=3, value=val)
                v_cell.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
                v_cell.alignment = Alignment(horizontal="right", vertical="center")

                # Highlight status
                if "В РОБОТІ" in str(val):
                    v_cell.fill = PatternFill(start_color=GREEN_BG, end_color=GREEN_BG, fill_type="solid")
                    v_cell.font = Font(name="Calibri", size=10, bold=True, color=GREEN_TXT)
                elif "ЗУПИНЕНО" in str(val):
                    v_cell.fill = PatternFill(start_color=RED_BG, end_color=RED_BG, fill_type="solid")
                    v_cell.font = Font(name="Calibri", size=10, bold=True, color=RED_TXT)

                is_even = (idx % 2 == 1)
                row_bg = ZEBRA_EVEN if is_even else ZEBRA_ODD
                if not v_cell.fill.start_color.rgb:
                    v_cell.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
                p_cell.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")

                border_thin = Side(border_style="thin", color=BORDER_COLOR)
                for col_i in range(1, 5):
                    ws_summary.cell(row=curr_row, column=col_i).border = Border(
                        top=border_thin, left=border_thin, right=border_thin, bottom=border_thin
                    )

                ws_summary.row_dimensions[curr_row].height = 20
                curr_row += 1

            curr_row += 1

        auto_fit_columns(ws_summary, min_width=18, max_width=50)

        # -----------------------------------------------------------------
        # 2. DAILY REPORT SHEET (📅 Добовий звіт)
        # -----------------------------------------------------------------
        ws_daily = wb.create_sheet(title="📅 Добовий звіт")
        ws_daily.sheet_properties.tabColor = "0EA5E9"
        ws_daily.views.sheetView[0].showGridLines = True
        ws_daily.freeze_panes = "A2"
        ws_daily.row_dimensions[1].height = 28

        daily_headers = [
            "Дата", "День тижня", "Залишок на початок (л)", "Сесії роботи (старт – зупинка)",
            "К-ть запусків", "Загальний час (год)", "Витрата палива (л)", "Заправлено (л)",
            "Хто привіз паливо", "Номер чеку / накладної", "Вартість заправки (₴)",
            "Залишок на кінець (л)", "Обслуговування (ТО) та події", "Відповідальні особи"
        ]
        for col_idx, h in enumerate(daily_headers, start=1):
            cell = ws_daily.cell(row=1, column=col_idx)
            apply_header_style(cell, h)

        # Grouping all events by Kyiv date
        days_map = {}
        for r in runs:
            st_loc = utc_to_local(r.start_time)
            sp_loc = utc_to_local(r.stop_time)
            d = st_loc.date()
            if d not in days_map:
                days_map[d] = {"runs": [], "refuels": [], "maints": [], "audits": []}
            days_map[d]["runs"].append((st_loc, sp_loc, r))

        for f in refuels:
            ts_loc = utc_to_local(f.timestamp)
            d = ts_loc.date()
            if d not in days_map:
                days_map[d] = {"runs": [], "refuels": [], "maints": [], "audits": []}
            days_map[d]["refuels"].append((ts_loc, f))

        for m in maints:
            ts_loc = utc_to_local(m.timestamp)
            d = ts_loc.date()
            if d not in days_map:
                days_map[d] = {"runs": [], "refuels": [], "maints": [], "audits": []}
            days_map[d]["maints"].append((ts_loc, m))

        for a in audits:
            ts_loc = utc_to_local(a.timestamp)
            d = ts_loc.date()
            if d not in days_map:
                days_map[d] = {"runs": [], "refuels": [], "maints": [], "audits": []}
            days_map[d]["audits"].append((ts_loc, a))

        sorted_dates = sorted(days_map.keys(), reverse=True)

        tot_day_runs_count = 0
        tot_day_hours = 0.0
        tot_day_fuel_burned = 0.0
        tot_day_fuel_added = 0.0
        tot_day_fuel_cost = 0.0

        for row_idx, d in enumerate(sorted_dates, start=2):
            d_data = days_map[d]
            is_even = (row_idx % 2 == 0)

            d_runs = sorted(d_data["runs"], key=lambda x: x[0])
            d_refuels = sorted(d_data["refuels"], key=lambda x: x[0])
            d_maints = sorted(d_data["maints"], key=lambda x: x[0])
            d_audits = sorted(d_data["audits"], key=lambda x: x[0])

            # Session breakdown
            session_lines = []
            for i, (st, sp, r) in enumerate(d_runs, start=1):
                op = r.start_user_name or "—"
                if r.stop_user_name and r.stop_user_name != r.start_user_name:
                    op += f" / {r.stop_user_name}"
                session_lines.append(
                    f"#{i}: {st.strftime('%H:%M')} – {sp.strftime('%H:%M')} "
                    f"({r.duration_hours:.2f} мч, {r.fuel_consumed:.1f} л) [{op}]"
                )
            sessions_text = "\n".join(session_lines) if session_lines else "Не запускався"

            day_runs_count = len(d_runs)
            day_hours = sum(r.duration_hours for _, _, r in d_runs)
            day_burned = sum(r.fuel_consumed for _, _, r in d_runs)

            # Refuels breakdown
            day_added = sum(f.amount_liters for _, f in d_refuels)
            day_cost = sum((f.cost or 0.0) for _, f in d_refuels)

            deliv_names = [f.delivered_by.strip() for _, f in d_refuels if f.delivered_by and f.delivered_by.strip()]
            deliv_text = ", ".join(dict.fromkeys(deliv_names)) if deliv_names else ("—" if day_added == 0 else "Не вказано")

            rec_nums = [f.receipt_number.strip() for _, f in d_refuels if f.receipt_number and f.receipt_number.strip()]
            rec_text = ", ".join(dict.fromkeys(rec_nums)) if rec_nums else ("—" if day_added == 0 else "Без чеку")

            cost_text = f"{day_cost:,.2f} ₴".replace(",", " ") if day_cost > 0 else "—"

            # Fuel balance (start & end of day)
            day_fuel_events = []
            for st, _, r in d_runs:
                day_fuel_events.append((st, r.start_fuel, r.end_fuel))
            for ts, f in d_refuels:
                day_fuel_events.append((ts, f.fuel_before, f.fuel_after))
            day_fuel_events.sort(key=lambda x: x[0])

            if day_fuel_events:
                start_fuel = round(day_fuel_events[0][1], 1)
                end_fuel = round(day_fuel_events[-1][2], 1)
            else:
                start_fuel = round(gen.current_fuel, 1) if gen else 0.0
                end_fuel = start_fuel

            # Maintenance & other events
            maint_lines = []
            for _, m in d_maints:
                m_tag = "🟢 Головне ТО" if m.is_main else f"🔧 Проміжне ({m.title or m.maint_type})"
                c_part = f", {m.cost:,.0f} ₴".replace(",", " ") if m.cost else ""
                maint_lines.append(f"{m_tag}: {m.description} [{m.user_name or '—'}{c_part}]")

            for _, a in d_audits:
                maint_lines.append(f"⚠️ Скидання: {a.reason} [{a.user_name or '—'}]")

            maint_events_text = "\n".join(maint_lines) if maint_lines else "Без подій"

            # Responsible users
            users_set = set()
            for _, _, r in d_runs:
                if r.start_user_name: users_set.add(r.start_user_name)
                if r.stop_user_name: users_set.add(r.stop_user_name)
            for _, f in d_refuels:
                if f.user_name: users_set.add(f.user_name)
                if f.delivered_by: users_set.add(f.delivered_by)
            for _, m in d_maints:
                if m.user_name: users_set.add(m.user_name)
            for _, a in d_audits:
                if a.user_name: users_set.add(a.user_name)
            clean_users = sorted(u for u in users_set if u and u not in ["—", "Не вказано"])
            resp_users_text = ", ".join(clean_users) if clean_users else "—"

            num_lines = max(len(session_lines), len(maint_lines), 1)
            ws_daily.row_dimensions[row_idx].height = max(24, num_lines * 19)

            apply_data_cell(ws_daily.cell(row=row_idx, column=1), d.strftime("%d.%m.%Y"), align="center", is_even=is_even)
            apply_data_cell(ws_daily.cell(row=row_idx, column=2), UKR_WEEKDAYS[d.weekday()], align="center", is_even=is_even)
            apply_data_cell(ws_daily.cell(row=row_idx, column=3), start_fuel, align="right", is_even=is_even, bold=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=4), sessions_text, align="left", is_even=is_even, wrap_text=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=5), day_runs_count, align="center", is_even=is_even)
            apply_data_cell(ws_daily.cell(row=row_idx, column=6), round(day_hours, 2), align="right", is_even=is_even, bold=True, text_color=BLUE_TXT)
            apply_data_cell(ws_daily.cell(row=row_idx, column=7), round(day_burned, 2), align="right", is_even=is_even, bold=True, text_color=RED_TXT)
            apply_data_cell(ws_daily.cell(row=row_idx, column=8), round(day_added, 1), align="right", is_even=is_even, bold=True, text_color=GREEN_TXT)
            apply_data_cell(ws_daily.cell(row=row_idx, column=9), deliv_text, align="left", is_even=is_even, wrap_text=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=10), rec_text, align="center", is_even=is_even, wrap_text=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=11), cost_text, align="right", is_even=is_even, bold=bool(day_cost > 0))
            apply_data_cell(ws_daily.cell(row=row_idx, column=12), end_fuel, align="right", is_even=is_even, bold=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=13), maint_events_text, align="left", is_even=is_even, wrap_text=True)
            apply_data_cell(ws_daily.cell(row=row_idx, column=14), resp_users_text, align="left", is_even=is_even, wrap_text=True)

            tot_day_runs_count += day_runs_count
            tot_day_hours += day_hours
            tot_day_fuel_burned += day_burned
            tot_day_fuel_added += day_added
            tot_day_fuel_cost += day_cost

        tot_row_daily = len(sorted_dates) + 2
        apply_total_row(ws_daily, tot_row_daily, len(daily_headers), label="Всього за період:", sums={
            5: tot_day_runs_count,
            6: round(tot_day_hours, 2),
            7: round(tot_day_fuel_burned, 2),
            8: round(tot_day_fuel_added, 1),
            11: f"{tot_day_fuel_cost:,.2f} ₴".replace(",", " ")
        })

        if len(sorted_dates) > 0:
            ws_daily.auto_filter.ref = f"A1:{get_column_letter(len(daily_headers))}{tot_row_daily - 1}"
        auto_fit_columns(ws_daily, min_width=12, max_width=50)

        # -----------------------------------------------------------------
        # 3. RUNS SHEET (⏱ Журнал запусків)
        # -----------------------------------------------------------------
        ws_runs = wb.create_sheet(title="⏱ Журнал запусків")
        ws_runs.sheet_properties.tabColor = "10B981"
        ws_runs.views.sheetView[0].showGridLines = True
        ws_runs.freeze_panes = "A2"
        ws_runs.row_dimensions[1].height = 28

        run_headers = [
            "№", "Час запуску (Київ)", "Час зупинки (Київ)", "Тривалість (год)",
            "Витрата (л)", "Норма (л/год)", "Бак до (л)", "Бак після (л)",
            "Мотогодини (мч)", "Запустив", "Зупинив", "Примітки"
        ]
        for col_idx, h in enumerate(run_headers, start=1):
            cell = ws_runs.cell(row=1, column=col_idx)
            apply_header_style(cell, h)

        for row_idx, r in enumerate(runs, start=2):
            ws_runs.row_dimensions[row_idx].height = 21
            is_even = (row_idx % 2 == 0)

            apply_data_cell(ws_runs.cell(row=row_idx, column=1), r.id, align="center", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=2), format_dt(r.start_time), align="center", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=3), format_dt(r.stop_time), align="center", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=4), round(r.duration_hours, 2), align="right", is_even=is_even, bold=True, text_color=BLUE_TXT)
            apply_data_cell(ws_runs.cell(row=row_idx, column=5), round(r.fuel_consumed, 2), align="right", is_even=is_even, bold=True, text_color=RED_TXT)
            apply_data_cell(ws_runs.cell(row=row_idx, column=6), round(r.fuel_rate, 2), align="right", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=7), round(r.start_fuel, 1), align="right", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=8), round(r.end_fuel, 1), align="right", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=9), round(r.total_hours_after, 2), align="right", is_even=is_even, bold=True)
            apply_data_cell(ws_runs.cell(row=row_idx, column=10), r.start_user_name or "—", align="left", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=11), r.stop_user_name or "—", align="left", is_even=is_even)
            apply_data_cell(ws_runs.cell(row=row_idx, column=12), r.notes or "", align="left", is_even=is_even)

        tot_row_runs = len(runs) + 2
        apply_total_row(ws_runs, tot_row_runs, len(run_headers), label="Всього за період:", sums={
            4: round(total_hours_ran, 2),
            5: round(total_fuel_burned, 2)
        })

        if len(runs) > 0:
            ws_runs.auto_filter.ref = f"A1:{get_column_letter(len(run_headers))}{tot_row_runs - 1}"
        auto_fit_columns(ws_runs, min_width=12, max_width=40)

        # -----------------------------------------------------------------
        # 4. REFUELS SHEET (⛽ Заправки)
        # -----------------------------------------------------------------
        ws_fuel = wb.create_sheet(title="⛽ Заправки")
        ws_fuel.sheet_properties.tabColor = "F59E0B"
        ws_fuel.views.sheetView[0].showGridLines = True
        ws_fuel.freeze_panes = "A2"
        ws_fuel.row_dimensions[1].height = 28

        fuel_headers = [
            "№", "Дата і час (Київ)", "Заправлено (л)", "Бак до (л)",
            "Бак після (л)", "Вартість (₴)", "Хто заправив",
            "Хто привіз (Прізвище)", "Номер чеку / накладної", "Примітки / АЗС"
        ]
        for col_idx, h in enumerate(fuel_headers, start=1):
            cell = ws_fuel.cell(row=1, column=col_idx)
            apply_header_style(cell, h)

        for row_idx, f in enumerate(refuels, start=2):
            ws_fuel.row_dimensions[row_idx].height = 21
            is_even = (row_idx % 2 == 0)

            apply_data_cell(ws_fuel.cell(row=row_idx, column=1), f.id, align="center", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=2), format_dt(f.timestamp), align="center", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=3), round(f.amount_liters, 1), align="right", is_even=is_even, bold=True, text_color=GREEN_TXT)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=4), round(f.fuel_before, 1), align="right", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=5), round(f.fuel_after, 1), align="right", is_even=is_even)
            cost_str = f"{f.cost:.2f} ₴" if f.cost else "—"
            apply_data_cell(ws_fuel.cell(row=row_idx, column=6), cost_str, align="right", is_even=is_even, bold=bool(f.cost))
            apply_data_cell(ws_fuel.cell(row=row_idx, column=7), f.user_name or "—", align="left", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=8), f.delivered_by or "—", align="left", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=9), f.receipt_number or "—", align="center", is_even=is_even)
            apply_data_cell(ws_fuel.cell(row=row_idx, column=10), f.notes or "", align="left", is_even=is_even)

        tot_row_fuel = len(refuels) + 2
        apply_total_row(ws_fuel, tot_row_fuel, len(fuel_headers), label="Всього заправлено:", sums={
            3: round(total_fuel_added, 1),
            6: f"{total_fuel_cost:,.2f} ₴".replace(",", " ")
        })

        if len(refuels) > 0:
            ws_fuel.auto_filter.ref = f"A1:{get_column_letter(len(fuel_headers))}{tot_row_fuel - 1}"
        auto_fit_columns(ws_fuel, min_width=13, max_width=40)

        # -----------------------------------------------------------------
        # 5. MAINTENANCE SHEET (🔧 Обслуговування)
        # -----------------------------------------------------------------
        ws_maint = wb.create_sheet(title="🔧 Обслуговування (ТО)")
        ws_maint.sheet_properties.tabColor = "6366F1"
        ws_maint.views.sheetView[0].showGridLines = True
        ws_maint.freeze_panes = "A2"
        ws_maint.row_dimensions[1].height = 28

        maint_headers = [
            "№", "Дата і час (Київ)", "Категорія ТО", "Мотогодини ТО",
            "Наступне ТО (мч)", "Опис виконаних робіт", "Замінені деталі / розхідники",
            "Вартість (₴)", "Хто виконав"
        ]
        for col_idx, h in enumerate(maint_headers, start=1):
            cell = ws_maint.cell(row=1, column=col_idx)
            apply_header_style(cell, h)

        for row_idx, m in enumerate(maints, start=2):
            ws_maint.row_dimensions[row_idx].height = 21
            is_even = (row_idx % 2 == 0)

            is_main = getattr(m, "is_main", True)
            type_label = "🟢 ГОЛОВНЕ (Олива)" if is_main else f"🔧 ПРОМІЖНЕ: {m.title or m.maint_type or 'Деталі'}"
            type_bg = GREEN_BG if is_main else BLUE_BG
            type_txt = GREEN_TXT if is_main else BLUE_TXT

            apply_data_cell(ws_maint.cell(row=row_idx, column=1), m.id, align="center", is_even=is_even)
            apply_data_cell(ws_maint.cell(row=row_idx, column=2), format_dt(m.timestamp), align="center", is_even=is_even)
            apply_data_cell(ws_maint.cell(row=row_idx, column=3), type_label, align="center", is_even=is_even, bold=True, text_color=type_txt, bg_color=type_bg)
            apply_data_cell(ws_maint.cell(row=row_idx, column=4), round(m.hours_at_maintenance, 1), align="right", is_even=is_even, bold=True)
            apply_data_cell(ws_maint.cell(row=row_idx, column=5), round(m.next_maintenance_hours, 1) if is_main else "—", align="right", is_even=is_even)
            apply_data_cell(ws_maint.cell(row=row_idx, column=6), m.description, align="left", is_even=is_even)
            apply_data_cell(ws_maint.cell(row=row_idx, column=7), m.parts_replaced or "", align="left", is_even=is_even)
            m_cost_str = f"{m.cost:.2f} ₴" if m.cost else "—"
            apply_data_cell(ws_maint.cell(row=row_idx, column=8), m_cost_str, align="right", is_even=is_even, bold=bool(m.cost))
            apply_data_cell(ws_maint.cell(row=row_idx, column=9), m.user_name or "—", align="left", is_even=is_even)

        tot_row_maint = len(maints) + 2
        apply_total_row(ws_maint, tot_row_maint, len(maint_headers), label="Всього витрат на ТО:", sums={
            8: f"{total_maint_cost:,.2f} ₴".replace(",", " ")
        })

        if len(maints) > 0:
            ws_maint.auto_filter.ref = f"A1:{get_column_letter(len(maint_headers))}{tot_row_maint - 1}"
        auto_fit_columns(ws_maint, min_width=14, max_width=45)

        # -----------------------------------------------------------------
        # 6. AUDIT RESET SHEET (🛡 Журнал аудиту)
        # -----------------------------------------------------------------
        ws_audit = wb.create_sheet(title="🛡 Журнал аудиту")
        ws_audit.sheet_properties.tabColor = "EF4444"
        ws_audit.views.sheetView[0].showGridLines = True
        ws_audit.freeze_panes = "A2"
        ws_audit.row_dimensions[1].height = 28

        audit_headers = [
            "№", "Дата і час (Київ)", "Тип скидання", "Обов'язкова причина",
            "Хто виконав", "Деталі операції та попередній стан"
        ]
        for col_idx, h in enumerate(audit_headers, start=1):
            cell = ws_audit.cell(row=1, column=col_idx)
            apply_header_style(cell, h)

        reset_labels = {
            "all": "💥 ПОВНЕ СКИДАННЯ",
            "fuel_zero": "⛽ ОБНУЛЕННЯ БАКА",
            "hours_zero": "⏱ ОБНУЛЕННЯ МОТОГОДИН",
            "maint_main": "🛠 СКИДАННЯ ГОЛОВНОГО ТО",
            "maint_intermediate": "🔧 СКИДАННЯ ПРОМІЖНОГО ТО",
        }

        for row_idx, a in enumerate(audits, start=2):
            ws_audit.row_dimensions[row_idx].height = 21
            is_even = (row_idx % 2 == 0)
            a_label = reset_labels.get(a.reset_type, a.reset_type)

            apply_data_cell(ws_audit.cell(row=row_idx, column=1), a.id, align="center", is_even=is_even)
            apply_data_cell(ws_audit.cell(row=row_idx, column=2), format_dt(a.timestamp), align="center", is_even=is_even)
            apply_data_cell(ws_audit.cell(row=row_idx, column=3), a_label, align="center", is_even=is_even, bold=True, text_color=RED_TXT, bg_color=RED_BG)
            apply_data_cell(ws_audit.cell(row=row_idx, column=4), a.reason, align="left", is_even=is_even, bold=True)
            apply_data_cell(ws_audit.cell(row=row_idx, column=5), a.user_name or "—", align="left", is_even=is_even)
            apply_data_cell(ws_audit.cell(row=row_idx, column=6), a.details or "", align="left", is_even=is_even)

        tot_row_audit = len(audits) + 2
        apply_total_row(ws_audit, tot_row_audit, len(audit_headers), label="Всього записів аудиту:", sums={
            3: f"{len(audits)} дій"
        })

        if len(audits) > 0:
            ws_audit.auto_filter.ref = f"A1:{get_column_letter(len(audit_headers))}{tot_row_audit - 1}"
        auto_fit_columns(ws_audit, min_width=14, max_width=55)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output
